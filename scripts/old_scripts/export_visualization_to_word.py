import os
from docx import Document
from docx.shared import Inches
from glob import glob

# 配置
VIS_BASE = r"e:/baozi/deeplabv3plus/Inspect/VisualizeCompare"
OUT_SAMPLE_DOCX = r"e:/baozi/deeplabv3plus/Inspect/VisualizeCompare/by_sample_visualization.docx"
OUT_EPOCH_DOCX = r"e:/baozi/deeplabv3plus/Inspect/VisualizeCompare/by_epoch_visualization.docx"

# 1. 按样本为主
sample_root = os.path.join(VIS_BASE, "by_sample")
sample_names = [d for d in os.listdir(sample_root) if os.path.isdir(os.path.join(sample_root, d))]
epoch_tags = set()
sample_dict = {}
for sample in sample_names:
    folder = os.path.join(sample_root, sample)
    pred_imgs = {}
    for f in os.listdir(folder):
        if f.endswith("_img.jpg"):
            img_path = os.path.abspath(os.path.join(folder, f))
        elif f.endswith("_gt_color.png"):
            gt_color = os.path.abspath(os.path.join(folder, f))
        elif f.endswith("_gt_gray.png"):
            gt_gray = os.path.abspath(os.path.join(folder, f))
        elif f.startswith("ep") and f.endswith("_pred.png"):
            tag = f.split("_")[0]
            pred_imgs[tag] = os.path.abspath(os.path.join(folder, f))
            epoch_tags.add(tag)
    sample_dict[sample] = {
        "img": img_path,
        "gt_color": gt_color,
        "gt_gray": gt_gray,
        **pred_imgs
    }
epoch_tags = sorted(list(epoch_tags), key=lambda x: int(x[2:]))

doc = Document()
doc.add_heading('按样本可视化对比', 0)

cols = 3 + len(epoch_tags)
table = doc.add_table(rows=1, cols=cols)
hdr = table.rows[0].cells
hdr[0].text = '样本名'
hdr[1].text = '原图'
hdr[2].text = '标签(彩色)'
for i, ep in enumerate(epoch_tags):
    hdr[3+i].text = ep + '_预测'

for sample, info in sample_dict.items():
    row = table.add_row().cells
    row[0].text = sample
    row[1].add_paragraph().add_run().add_picture(info["img"], width=Inches(1.0))
    row[2].add_paragraph().add_run().add_picture(info["gt_color"], width=Inches(1.0))
    for i, ep in enumerate(epoch_tags):
        if ep in info:
            row[3+i].add_paragraph().add_run().add_picture(info[ep], width=Inches(1.0))
        else:
            row[3+i].text = "-"
doc.save(OUT_SAMPLE_DOCX)
print(f"已生成按样本可视化Word: {OUT_SAMPLE_DOCX}")

# 2. 按轮次为主
epoch_folders = [d for d in os.listdir(VIS_BASE) if d.startswith("val_compare_epoch")]
epoch_folders = sorted(epoch_folders, key=lambda x: int(x.split("epoch")[-1]))
sample_tags = set()
epoch_dict = {}
for epf in epoch_folders:
    folder = os.path.join(VIS_BASE, epf)
    pred_imgs = {}
    for f in os.listdir(folder):
        if f.endswith(".png"):
            sample = f.replace(".png", "")
            pred_imgs[sample] = os.path.abspath(os.path.join(folder, f))
            sample_tags.add(sample)
    epoch_dict[epf] = pred_imgs
sample_tags = sorted(list(sample_tags))

doc2 = Document()
doc2.add_heading('按轮次可视化对比', 0)
cols2 = 1 + len(sample_tags)
table2 = doc2.add_table(rows=1, cols=cols2)
hdr2 = table2.rows[0].cells
hdr2[0].text = '轮次'
for i, s in enumerate(sample_tags):
    hdr2[1+i].text = s
for epf, pred_imgs in epoch_dict.items():
    row = table2.add_row().cells
    row[0].text = epf
    for i, s in enumerate(sample_tags):
        if s in pred_imgs:
            row[1+i].add_paragraph().add_run().add_picture(pred_imgs[s], width=Inches(1.0))
        else:
            row[1+i].text = "-"
doc2.save(OUT_EPOCH_DOCX)
print(f"已生成按轮次可视化Word: {OUT_EPOCH_DOCX}")
